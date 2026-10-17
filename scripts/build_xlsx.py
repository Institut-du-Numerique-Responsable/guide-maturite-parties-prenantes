"""Build the Excel workbooks from the Markdown chapters.

For each language, two workbooks:
- the evaluation grid (criteria, weights, scoring and a per-chapter radar),
  kept by the evaluator;
- the questionnaire (open questions and rationales only), sent to the
  stakeholders, since the guide advises not to disclose the criteria.

Usage: python scripts/build_xlsx.py OUT_DIR
"""

from __future__ import annotations

import pathlib
import sys

from openpyxl import Workbook
from openpyxl.chart import RadarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import guide

BRAND = "00668B"
LIGHT = "E6F1F5"
VERSION = "1.2"
URL = "https://institutnr.org/guide-maturite-parties-prenantes"
REPO = "https://github.com/Institut-du-Numerique-Responsable/guide-maturite-parties-prenantes"

TEXT = {
    "fr": dict(
        grid_file="grille-evaluation-fr.xlsx",
        questionnaire_file="questionnaire-fr.xlsx",
        title="Guide d'évaluation et d'influence de la maturité numérique responsable des parties prenantes",
        grid_title="Grille d'évaluation",
        summary="Synthèse",
        howto="Mode d'emploi",
        questionnaire="Questionnaire",
        cols=["Chapitre", "N°", "Question ouverte", "Critère d'évaluation", "Priorité", "Pondération",
              "Niveau atteint", "Points obtenus", "Éléments de preuve fournis", "Commentaire de l'évaluateur"],
        qcols=["N°", "Question ouverte", "Raison d'être", "Réponse de la partie prenante", "Pièces jointes / références (page)"],
        scols=["Chapitre", "Points possibles", "Points obtenus", "Maturité"],
        stakeholder="Partie prenante évaluée :",
        evaluator="Évaluateur(s) :",
        date="Date :",
        total="Total",
        radar="Maturité par famille de critères",
        no_rationale="Voir la raison d'être de la question équivalente au chapitre 5.",
        howto_lines=[
            "Cette grille est destinée à l'évaluateur. Ne la transmettez pas aux parties prenantes : envoyez-leur le questionnaire, qui ne contient que les questions ouvertes et leur raison d'être.",
            "1. Renseignez l'organisation évaluée, les évaluateurs et la date en haut de l'onglet « Grille d'évaluation ».",
            "2. Pour chaque critère, choisissez le niveau atteint dans la liste : 0 %, 25 %, 50 %, 75 % ou 100 %. Les points obtenus sont calculés automatiquement (pondération × niveau).",
            "3. Notez les éléments de preuve fournis (certificats, rapports, numéros de page) : seules les informations fournies dans le questionnaire sont prises en compte.",
            "4. L'onglet « Synthèse » calcule la maturité par famille de critères (100 points par chapitre) et trace le radar à présenter lors du débriefing.",
            "Une maturité de 100 % n'est pas réaliste : l'objectif est d'identifier le principal point fort à sécuriser et les axes d'amélioration justifiant un plan d'action.",
            "Vous pouvez ne retenir qu'une partie des chapitres et adapter pondérations ou priorités ; P1 = prioritaire, P2 = recommandé, P3 = pour aller plus loin.",
        ],
        cite=("Source : Institut du Numérique Responsable (INR) / Institute for Sustainable IT (ISIT), Guide d'évaluation et "
              f"d'influence de la maturité numérique responsable des parties prenantes, version {VERSION}. Licence CC BY-SA 4.0 : "
              f"toute réutilisation doit citer l'INR. {URL}"),
        q_intro="Merci de répondre à chaque question en synthétisant les informations disponibles (évitez les simples liens vers un rapport ; indiquez le numéro de page). Seules les informations fournies dans ce questionnaire seront prises en compte.",
    ),
    "en": dict(
        grid_file="evaluation-grid-en.xlsx",
        questionnaire_file="questionnaire-en.xlsx",
        title="Guide to evaluating and influencing the maturity of stakeholders in Sustainable IT",
        grid_title="Evaluation grid",
        summary="Summary",
        howto="How to use",
        questionnaire="Questionnaire",
        cols=["Chapter", "No.", "Open question", "Evaluation criterion", "Priority", "Weighting",
              "Level reached", "Points obtained", "Evidence provided", "Evaluator comment"],
        qcols=["No.", "Open question", "Rationale", "Stakeholder answer", "Attachments / references (page)"],
        scols=["Chapter", "Possible points", "Points obtained", "Maturity"],
        stakeholder="Evaluated stakeholder:",
        evaluator="Evaluator(s):",
        date="Date:",
        total="Total",
        radar="Maturity per family of criteria",
        no_rationale="See the rationale of the equivalent question in chapter 5.",
        howto_lines=[
            "This grid is meant for the evaluator. Do not send it to stakeholders: send them the questionnaire, which only contains the open questions and their rationales.",
            "1. Fill in the evaluated organisation, the evaluators and the date at the top of the \"Evaluation grid\" sheet.",
            "2. For each criterion, pick the level reached from the list: 0%, 25%, 50%, 75% or 100%. Points obtained are computed automatically (weighting × level).",
            "3. Record the evidence provided (certificates, reports, page numbers): only information provided in the questionnaire is taken into account.",
            "4. The \"Summary\" sheet computes the maturity per family of criteria (100 points per chapter) and draws the radar chart to present during the debriefing.",
            "A 100% maturity is not realistic: the goal is to identify the main strength to secure and the improvement areas that justify an action plan.",
            "You may keep only some chapters and adapt weightings or priorities; P1 = priority, P2 = recommended, P3 = going further.",
        ],
        cite=("Source: Institut du Numérique Responsable (INR) / Institute for Sustainable IT (ISIT), Guide to evaluating and "
              f"influencing the maturity of stakeholders in Sustainable IT, version {VERSION}. Licence CC BY-SA 4.0: any reuse "
              f"must cite the INR. {URL}"),
        q_intro="Please answer each question by synthesising the available information (avoid bare links to a report; give the page number). Only information provided in this questionnaire will be taken into account.",
    ),
}

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=BRAND)
BAND_FILL = PatternFill("solid", fgColor=LIGHT)
INPUT_FILL = PatternFill("solid", fgColor="FFF8E1")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def header(ws, row: int, labels: list[str], widths: list[int]) -> None:
    for i, (label, width) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.alignment, c.border = HEAD_FONT, HEAD_FILL, CENTER, BOX
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def title_block(ws, t: dict, subtitle: str) -> None:
    ws["A1"] = t["title"]
    ws["A1"].font = Font(bold=True, size=14, color=BRAND)
    ws["A2"] = subtitle
    ws["A2"].font = Font(bold=True, size=12)


def footer(ws, row: int, t: dict) -> None:
    c = ws.cell(row=row, column=1, value=t["cite"])
    c.font = Font(italic=True, size=9, color="595959")


def set_properties(wb: Workbook, t: dict) -> None:
    wb.properties.creator = "Institut du Numérique Responsable (INR / ISIT)"
    wb.properties.title = t["title"]
    wb.properties.subject = t["cite"]
    wb.properties.keywords = "numérique responsable; sustainable IT; maturité; INR; ISIT; CC BY-SA 4.0"


def build_grid(lang: str, chapters: list[guide.Chapter], out: pathlib.Path) -> None:
    t = TEXT[lang]
    wb = Workbook()
    set_properties(wb, t)

    howto = wb.active
    howto.title = t["howto"]
    title_block(howto, t, t["howto"])
    howto.column_dimensions["A"].width = 120
    for i, line in enumerate(t["howto_lines"], start=4):
        howto.cell(row=i, column=1, value=line).alignment = WRAP
    footer(howto, 4 + len(t["howto_lines"]) + 1, t)

    ws = wb.create_sheet(t["grid_title"])
    title_block(ws, t, t["grid_title"])
    for r, label in enumerate((t["stakeholder"], t["evaluator"], t["date"]), start=3):
        ws.cell(row=r, column=3, value=label).font = Font(bold=True)
        ws.cell(row=r, column=4).fill = INPUT_FILL
    first = 7
    header(ws, first, t["cols"], [30, 6, 45, 70, 9, 11, 12, 11, 40, 40])
    levels = DataValidation(type="list", formula1='"0%,25%,50%,75%,100%"', allow_blank=True)
    ws.add_data_validation(levels)
    row = first + 1
    for ch in chapters:
        for qi, q in enumerate(ch.questions):
            for ci, c in enumerate(q.criteria):
                values = [ch.title, q.id, q.title, guide.plain(c.text), c.priority, c.weight, None,
                          f"=IF(G{row}=\"\",0,F{row}*G{row})", None, None]
                for col, v in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=col, value=v)
                    cell.alignment = CENTER if col in (2, 5, 6, 7, 8) else WRAP
                    cell.border = BOX
                    if qi % 2:
                        cell.fill = BAND_FILL
                for col in (7, 9, 10):
                    ws.cell(row=row, column=col).fill = INPUT_FILL
                ws.cell(row=row, column=7).number_format = "0%"
                ws.cell(row=row, column=8).number_format = "0.##"
                levels.add(ws.cell(row=row, column=7))
                row += 1
    last = row - 1
    ws.auto_filter.ref = f"A{first}:J{last}"
    ws.cell(row=row, column=5, value=t["total"]).font = Font(bold=True)
    ws.cell(row=row, column=6, value=f"=SUM(F{first + 1}:F{last})").font = Font(bold=True)
    ws.cell(row=row, column=8, value=f"=SUM(H{first + 1}:H{last})").font = Font(bold=True)
    footer(ws, row + 2, t)

    summary = wb.create_sheet(t["summary"], 1)
    title_block(summary, t, t["summary"])
    header(summary, 3, t["scols"], [55, 16, 16, 12])
    grid = "'" + t["grid_title"].replace("'", "''") + "'"
    for i, ch in enumerate(chapters, start=4):
        summary.cell(row=i, column=1, value=ch.title).alignment = WRAP
        summary.cell(row=i, column=2, value=f"=SUMIF({grid}!$A${first + 1}:$A${last},A{i},{grid}!$F${first + 1}:$F${last})")
        summary.cell(row=i, column=3, value=f"=SUMIF({grid}!$A${first + 1}:$A${last},A{i},{grid}!$H${first + 1}:$H${last})")
        summary.cell(row=i, column=4, value=f"=IF(B{i}=0,0,C{i}/B{i})").number_format = "0%"
        for col in range(1, 5):
            summary.cell(row=i, column=col).border = BOX
    end = 3 + len(chapters)
    summary.conditional_formatting.add(f"D4:D{end}", CellIsRule(operator="lessThan", formula=["0.25"], fill=PatternFill("solid", fgColor="F8D7DA")))
    summary.conditional_formatting.add(f"D4:D{end}", CellIsRule(operator="greaterThanOrEqual", formula=["0.6"], fill=PatternFill("solid", fgColor="D4EDDA")))
    chart = RadarChart()
    chart.type = "marker"
    chart.title = t["radar"]
    chart.style = 26
    chart.add_data(Reference(summary, min_col=4, min_row=3, max_row=end), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=4, max_row=end))
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1
    chart.y_axis.number_format = "0%"
    chart.height, chart.width = 11, 18
    summary.add_chart(chart, "F3")
    footer(summary, end + 2, t)

    wb.active = 1
    wb.save(out / t["grid_file"])


def build_questionnaire(lang: str, chapters: list[guide.Chapter], out: pathlib.Path) -> None:
    t = TEXT[lang]
    wb = Workbook()
    set_properties(wb, t)
    ws = wb.active
    ws.title = t["questionnaire"]
    title_block(ws, t, t["questionnaire"])
    ws["A3"] = t["q_intro"]
    ws["A3"].alignment = WRAP
    ws.merge_cells("A3:E3")
    ws.row_dimensions[3].height = 45
    header(ws, 5, t["qcols"], [7, 55, 70, 70, 30])
    row = 6
    for ch in chapters:
        c = ws.cell(row=row, column=1, value=ch.title)
        c.font = Font(bold=True, color=BRAND)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 1
        for q in ch.questions:
            values = [q.id, guide.plain(q.title), q.rationale or t["no_rationale"], None, None]
            for col, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=v)
                cell.alignment = CENTER if col == 1 else WRAP
                cell.border = BOX
            for col in (4, 5):
                ws.cell(row=row, column=col).fill = INPUT_FILL
            row += 1
    footer(ws, row + 1, t)
    wb.save(out / t["questionnaire_file"])


def main(argv: list[str]) -> int:
    if len(argv) != 1:
        print(__doc__)
        return 2
    out = pathlib.Path(argv[0])
    out.mkdir(parents=True, exist_ok=True)
    for lang in guide.LANGS:
        chapters = guide.load(lang)
        build_grid(lang, chapters, out)
        build_questionnaire(lang, chapters, out)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
